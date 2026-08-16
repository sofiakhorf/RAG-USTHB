import os
import logging
from typing import List, Optional
import pymupdf
from docx import Document

from core import Chunk, ChunkMetadata

logger = logging.getLogger(__name__)


class FallbackParser:
    """Parseur de secours (dégradation gracieuse) en cas d'échec d'un parseur spécialisé."""

    @staticmethod
    def parse_safe(file_path: str, family: str, document_title: Optional[str] = None) -> List[Chunk]:
        clean_title = document_title if document_title else os.path.basename(file_path)
        logger.warning(f"Bascule sur le parseur de secours pour : {file_path} (Famille {family})")

        extracted_lines: List[str] = []
        ext = os.path.splitext(file_path)[1].lower()

        try:
            if ext == ".pdf":
                with pymupdf.open(file_path) as doc:
                    for page in doc:
                        txt = page.get_text()
                        if txt and txt.strip():
                            extracted_lines.append(txt.strip())

            elif ext == ".docx":
                doc = Document(file_path)
                # Extraction des paragraphes
                for p in doc.paragraphs:
                    if p.text and p.text.strip():
                        extracted_lines.append(p.text.strip())
                # Extraction des tableaux
                for table in doc.tables:
                    for row in table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            extracted_lines.append(" | ".join(cells))

            elif ext in [".txt", ".csv", ".md", ".json", ".log", ".tsv"]:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
                    if content.strip():
                        extracted_lines.append(content.strip())

            else:
                extracted_lines.append(f"Fichier non textuel brut : {os.path.basename(file_path)}")

            extracted_text = "\n".join(extracted_lines)

        except Exception as e:
            logger.error(f"Échec critique de l'extraction fallback sur {file_path} : {e}")
            extracted_text = f"Échec total de l'extraction. Fichier hérité non lisible. Erreur: {str(e)}"

        return [Chunk(
            text=extracted_text.strip() or "Contenu vide",
            metadata=ChunkMetadata(
                document_name=clean_title,
                family_type=family,
                is_fallback=True
            )
        )]