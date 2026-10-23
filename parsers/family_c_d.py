# -*- coding: utf-8 -*-
"""
family_c_d.py — Parser Excel (familles C et D) pour le RAG de la Faculté d'Informatique.

Principe : on ne se base pas sur des noms de fichiers ni sur des positions de cellules codées en dur,
mais sur la STRUCTURE du classeur (lignes vides, cellules fusionnées, en-têtes à plusieurs niveaux).

    SheetView      lecture unique de la feuille + table des fusions (valeurs normalisées)
    analyze_block  un bloc (lignes contiguës non vides) est classé :
                     - "banner" : bandeau / en-tête de document (titre, année, version…)
                     - "text"   : liste de texte libre dans une seule colonne (sujets, notes…)
                     - "table"  : tableau, avec en-tête à 1..n niveaux (bandeaux de groupes + ligne de libellés)
    ExcelWorkbookParser
                     un tableau est rendu selon sa forme :
                     - grille "groupes de colonnes répétés" (salles x créneaux)  -> famille D, 1 chunk / soutenance
                     - liste d'étudiants (Matricule + Nom)                        -> famille C, 1 chunk / étudiant
                                                                                    + effectif + listes par groupe
                     - tableau plat                                               -> famille C, 1 chunk / ligne
                     - liste de texte libre                                       -> famille C-texte-libre

Chaque chunk est autonome : il porte son contexte (document, feuille, titres, section) dans le texte,
car c'est le texte qui est vectorisé.
"""
import os
import re
import logging
from collections import Counter, OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

# --- Imports alignés sur la nouvelle architecture modulaire ---
from core import Chunk, ChunkMetadata, BaseFamilyParser, ParserRegistry
from core.utils import AcademicMetadataExtractor
from parsers.fallback import FallbackParser  # noqa: F401  (import conservé tel quel)

logger = logging.getLogger("RAG_USTHB_EXCEL")


# =====================================================================
# 1. NORMALISATION DES VALEURS
# =====================================================================

_WS_RE = re.compile(r"[\s ​‎‏]+")
_FR_DAYS = ("lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche")
_FR_MONTHS = ("janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
              "septembre", "octobre", "novembre", "décembre")


def format_date_fr(d: date) -> str:
    """Date lisible ET recherchable : 'dimanche 1er juin 2025 (01/06/2025)'."""
    jour = "1er" if d.day == 1 else str(d.day)
    return f"{_FR_DAYS[d.weekday()]} {jour} {_FR_MONTHS[d.month - 1]} {d.year} ({d:%d/%m/%Y})"


def clean_value(v: Any) -> Optional[str]:
    """Valeur de cellule -> texte propre (espaces/retours à la ligne compactés), None si vide."""
    if v is None:
        return None
    if isinstance(v, bool):
        return "Oui" if v else "Non"
    if isinstance(v, datetime):
        base = format_date_fr(v.date())
        return base if (v.hour, v.minute) == (0, 0) else f"{base} à {v.hour}h{v.minute:02d}"
    if isinstance(v, date):
        return format_date_fr(v)
    if isinstance(v, time):
        return f"{v.hour}h{v.minute:02d}"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v != v:  # NaN
            return None
        return str(int(v)) if v.is_integer() else format(v, ".10g")
    text = _WS_RE.sub(" ", str(v)).strip()
    return text or None


_PEOPLE_SPLIT_RE = re.compile(r"\s*(?:-{2,}|\s+[–—]\s*|[–—]\s+)\s*")


def split_people(text: Optional[str]) -> List[str]:
    """'A, X  --  B, Y  -- ' -> ['A, X', 'B, Y'] (séparateurs '--' et vides supprimés)."""
    if not text:
        return []
    return [p.strip(" ,;") for p in _PEOPLE_SPLIT_RE.split(text) if p.strip(" ,;")]


_YEAR_PAIR_RE = re.compile(r"(?<!\d)(20\d{2})\s*[-/_.]?\s*(20\d{2})(?!\d)")


def year_pair(text: Optional[str]) -> Optional[str]:
    """'20232024' / '2023-2024' / '2023/2024' -> '2023-2024' (années consécutives uniquement)."""
    for m in _YEAR_PAIR_RE.finditer(text or ""):
        a, b = int(m.group(1)), int(m.group(2))
        if b == a + 1:
            return f"{a}-{b}"
    return None


def _tokens(text: str) -> set:
    return {t[:-1] if len(t) > 3 and t.endswith("s") else t for t in re.findall(r"\w+", text.lower())}


def dedupe_parts(parts: List[str]) -> List[str]:
    """Retire d'un en-tête les morceaux déjà couverts par les autres (ex. nom de feuille = titre)."""
    parts = [p for p in parts if p]
    keep = [True] * len(parts)
    covered: set = set()
    for i in sorted(range(len(parts)), key=lambda k: -len(_tokens(parts[k]))):
        tk = _tokens(parts[i])
        if tk and tk <= covered:
            keep[i] = False
        else:
            covered |= tk
    return [p for p, k in zip(parts, keep) if k]


def dedupe_consecutive(items: List[str]) -> List[str]:
    out: List[str] = []
    for it in items:
        if not out or out[-1] != it:
            out.append(it)
    return out


# --- Repères linguistiques des documents administratifs ---
NOTE_RE = re.compile(r"^\s*(?:(?:NB|N\.B\.?)\b[:.\-]?|(?:Note|Remarque|Attention|Important)\s*[:.\-])\s*", re.IGNORECASE)
SECTION_RE = re.compile(r"^Sujets?\s+(?:disponibles|licence|master|pfe|externes|acad|isil|gtr)\b.*:?$", re.IGNORECASE)
SOURCE_RE = re.compile(r"^Source\s*:\s*(.*)$", re.IGNORECASE)
VERSION_RE = re.compile(r"^\s*version\b", re.IGNORECASE)
INSTITUTION_RE = re.compile(
    r"^\s*(?:r[ée]publique|minist[èe]re|universit[ée]|facult[ée]|d[ée]partement|[ée]cole)\b", re.IGNORECASE)
BULLET_RE = re.compile(r"^[·•●*\-–—]\s*")
ITEM_NUMBER_RE = re.compile(r"^\d{1,2}\s*[-.)]\s+(?=[A-Za-zÀ-ÿ0-9])")
HEADER_SECTION_RE = re.compile(r"^(Sujets?)\s+(\S.*)$", re.IGNORECASE)
# Colonnes de suivi hebdomadaire (Présences / Sem 1 / C-TD-TP) : sans intérêt pour la recherche
ATTENDANCE_RE = re.compile(r"^(?:pr[ée]sences?|absences?|assiduit[ée]|sem(?:aine)?\.?\s*\d+)$", re.IGNORECASE)
TIME_RANGE_RE = re.compile(r"^\d{1,2}\s*h\s*\d{0,2}\s*[-–]\s*\d{1,2}\s*h\s*\d{0,2}$", re.IGNORECASE)


# =====================================================================
# 2. VUE SUR UNE FEUILLE (valeurs + fusions)
# =====================================================================

@dataclass
class Segment:
    """Cellule non vide d'une ligne, avec l'étendue de sa fusion éventuelle."""
    c0: int
    c1: int
    vspan: int
    text: str

    @property
    def hspan(self) -> int:
        return self.c1 - self.c0 + 1


class SheetView:
    """Feuille chargée une seule fois en mémoire (indices 1-based, comme Excel)."""

    MAX_MERGE_AREA = 50_000

    def __init__(self, ws):
        self.title = ws.title
        self._orig = [list(row) for row in ws.iter_rows(values_only=True)]
        self._cells = [[clean_value(v) for v in row] for row in self._orig]
        self.nrows = self.ncols = 0
        for r, row in enumerate(self._cells, 1):
            for c, v in enumerate(row, 1):
                if v is not None:
                    self.nrows, self.ncols = max(self.nrows, r), max(self.ncols, c)
        self.merge_of: Dict[Tuple[int, int], Tuple[int, int, int, int]] = {}
        for rng in ws.merged_cells.ranges:
            r1, c1 = min(rng.max_row, self.nrows), min(rng.max_col, self.ncols)
            if (r1 - rng.min_row + 1) * (c1 - rng.min_col + 1) > self.MAX_MERGE_AREA:
                continue
            box = (rng.min_row, rng.min_col, r1, c1)
            for r in range(rng.min_row, r1 + 1):
                for c in range(rng.min_col, c1 + 1):
                    self.merge_of[(r, c)] = box

    # --- accès cellule ---
    def raw(self, r: int, c: int) -> Optional[str]:
        if 1 <= r <= len(self._cells) and 1 <= c <= len(self._cells[r - 1]):
            return self._cells[r - 1][c - 1]
        return None

    def orig(self, r: int, c: int) -> Any:
        if 1 <= r <= len(self._orig) and 1 <= c <= len(self._orig[r - 1]):
            return self._orig[r - 1][c - 1]
        return None

    def filled(self, r: int, c: int) -> Optional[str]:
        """Valeur avec propagation des fusions (utile pour les en-têtes)."""
        v = self.raw(r, c)
        if v is None and (r, c) in self.merge_of:
            b = self.merge_of[(r, c)]
            return self.raw(b[0], b[1])
        return v

    def body_value(self, r: int, c: int) -> Optional[str]:
        """Valeur d'une cellule de données : une fusion verticale (ex. la date d'une journée) se propage
        vers le bas dans sa colonne d'ancrage ; une fusion horizontale n'est pas dupliquée."""
        v = self.raw(r, c)
        if v is None and (r, c) in self.merge_of:
            b = self.merge_of[(r, c)]
            if c == b[1]:
                return self.raw(b[0], b[1])
        return v

    def segments(self, r: int) -> List[Segment]:
        out = []
        for c in range(1, (len(self._cells[r - 1]) if 1 <= r <= len(self._cells) else 0) + 1):
            v = self.raw(r, c)
            if v is None:
                continue
            b = self.merge_of.get((r, c), (r, c, r, c))
            out.append(Segment(c0=c, c1=b[3], vspan=b[2] - b[0] + 1, text=v))
        return out

    # --- segmentation en blocs ---
    def _is_blank(self, r: int) -> bool:
        return all(self.filled(r, c) is None for c in range(1, self.ncols + 1))

    def blocks(self, start_line: int = 0) -> List[Tuple[int, int]]:
        """Blocs de lignes contiguës non vides. `start_line` : première ligne à prendre en compte."""
        found, start = [], None
        for r in range(1, self.nrows + 1):
            if self._is_blank(r):
                if start is not None:
                    found.append((start, r - 1))
                    start = None
            elif start is None:
                start = r
        if start is not None:
            found.append((start, self.nrows))
        return [(max(a, start_line), b) for a, b in found if b >= start_line]


# =====================================================================
# 3. ANALYSE D'UN BLOC (bandeau / texte libre / tableau)
# =====================================================================

@dataclass
class Column:
    idx: int                       # colonne Excel (1-based)
    path: List[str]                # [libellé(s) de groupe..., libellé de colonne]
    leaf: Optional[str] = None     # libellé de colonne seul (None si colonne sans titre)


@dataclass
class BlockLayout:
    kind: str                      # "banner" | "text" | "table"
    r0: int
    r1: int
    titles: List[str] = field(default_factory=list)       # titres / mentions au-dessus du tableau
    columns: List[Column] = field(default_factory=list)
    leaf_row: int = 0
    table_r0: int = 0              # première ligne d'en-tête du tableau (hors titres)
    body_rows: List[int] = field(default_factory=list)
    text_col: int = 0


def analyze_block(view: SheetView, r0: int, r1: int) -> BlockLayout:
    rows = {r: view.segments(r) for r in range(r0, r1 + 1)}
    nonempty = [s for segs in rows.values() for s in segs]
    if not nonempty:
        return BlockLayout("banner", r0, r1)

    # Bandeau : toutes les lignes sont une cellule unique fusionnée sur plusieurs colonnes
    if all(len(segs) == 1 and segs[0].hspan >= 2 for segs in rows.values() if segs):
        return BlockLayout("banner", r0, r1, titles=[s.text for s in nonempty])
    # Texte libre : une seule colonne, aucune fusion
    if len({s.c0 for s in nonempty}) == 1 and all(s.hspan == 1 for s in nonempty):
        return BlockLayout("text", r0, r1, text_col=nonempty[0].c0)

    block_c0 = min(s.c0 for s in nonempty)
    block_c1 = max(s.c1 for s in nonempty)

    # Lignes "bandeau" en tête : titres et libellés de groupes de colonnes (cellules fusionnées)
    band_rows, r = [], r0
    while r <= r1:
        segs = rows[r]
        if not segs:
            r += 1
            continue
        if len(segs) == 1 or all(s.hspan >= 2 or s.vspan >= 2 for s in segs):
            band_rows.append(r)
            r += 1
        else:
            break
    leaf_row = r
    if leaf_row > r1:  # que des bandeaux, pas de tableau
        return BlockLayout("banner", r0, r1, titles=[s.text for s in nonempty])

    titles: List[str] = []
    levels: List[List[Segment]] = []
    for br in band_rows:
        segs, nxt, level = rows[br], br + 1, []
        for s in segs:
            if s.vspan >= 2:      # libellé de colonne fusionné sur plusieurs lignes d'en-tête (ex. JOURNEE)
                continue
            full_width = len(segs) == 1 and s.c0 == block_c0 and (s.hspan >= 0.5 * (block_c1 - block_c0 + 1) or not has_children)
            has_children = nxt <= r1 and any(view.filled(nxt, c) for c in range(s.c0, s.c1 + 1))
            if full_width or not has_children or (s.hspan == 1 and len(segs) == 1):
                titles.append(s.text)
            else:
                level.append(s)
        if level:
            levels.append(level)

    columns: List[Column] = []
    for c in range(block_c0, block_c1 + 1):
        path = []
        for level in levels:
            for s in level:
                if s.c0 <= c <= s.c1:
                    path.append(s.text)
                    break
        leaf = view.filled(leaf_row, c)
        if leaf:
            path.append(leaf)
        columns.append(Column(idx=c, path=dedupe_consecutive(path), leaf=leaf))

    body = [r for r in range(leaf_row + 1, r1 + 1)
            if any(view.body_value(r, c) is not None for c in range(block_c0, block_c1 + 1))]
    first_level_row = next((br for br in band_rows
                            if any(s.vspan < 2 and s.text not in titles for s in rows[br])), leaf_row)
    return BlockLayout("table", r0, r1, titles=titles, columns=columns, leaf_row=leaf_row,
                       table_r0=min(first_level_row, leaf_row), body_rows=body)


# =====================================================================
# 4. CONTEXTE DE FEUILLE ET RENDU
# =====================================================================

@dataclass
class SheetContext:
    doc_name: str                           # nom de fichier / titre fourni (metadata.document_name)
    doc_title: str                          # même chose, nettoyé pour l'en-tête des chunks
    sheet_name: str
    sheet_display: Optional[str]            # None s'il n'y a qu'une feuille visible (le nom n'apporte rien)
    preamble: List[str] = field(default_factory=list)   # lignes d'en-tête utiles (année, intitulé de liste…)
    version: Optional[str] = None
    notes: List[str] = field(default_factory=list)
    file_meta: Dict[str, Any] = field(default_factory=dict)
    sheet_meta: Dict[str, Any] = field(default_factory=dict)
    pre_meta: Dict[str, Any] = field(default_factory=dict)
    planning_hint: bool = False


@dataclass
class TextState:
    section: Optional[str] = None
    source: Optional[str] = None


def split_items(view: SheetView, r: int, c: int) -> List[str]:
    """Une cellule de texte libre -> un ou plusieurs éléments.
    Découpage numéroté (1- 2- 3-…) seulement si la cellule COMMENCE par une liste numérotée ; sinon une
    puce suivie de sous-points reste un seul élément (le sous-point garde le contexte de son en-tête)."""
    orig = view.orig(r, c)
    if not isinstance(orig, str):
        t = view.raw(r, c)
        return [t] if t else []
    lines = [ln.strip() for ln in orig.replace("\r", "").split("\n") if ln.strip()]
    if not lines:
        return []
    marks = [i for i, ln in enumerate(lines) if ITEM_NUMBER_RE.match(ln)]
    if len(marks) >= 2 and marks[0] == 0:
        items: List[str] = []
        for ln in lines:
            if ITEM_NUMBER_RE.match(ln):
                items.append(ITEM_NUMBER_RE.sub("", ln, count=1))
            else:
                items[-1] += " " + ln
        return [t for t in (clean_value(i) for i in items) if t]
    text = ITEM_NUMBER_RE.sub("", BULLET_RE.sub("", view.raw(r, c) or ""), count=1).strip()
    return [text] if text else []


def _natural_key(v: str):
    return (0, int(v), "") if v.isdigit() else (1, 0, v.lower())


@ParserRegistry.register("EXCEL")
class ExcelWorkbookParser(BaseFamilyParser):
    """Familles C (tableaux, listes, texte libre) et D (plannings / grilles de soutenances)."""

    def __init__(self, config=None):
        super().__init__(config)

    # ------------------------------------------------------------------
    # Indices (le nom de fichier n'est qu'un indice, jamais une condition nécessaire)
    # ------------------------------------------------------------------
    @staticmethod
    def _is_planning_file(filename: str) -> bool:
        """Indice 'planning' ET 'soutenance/pfe' dans le nom du fichier (sert pour les tableaux plats)."""
        has_planning = bool(re.search(r"planning|emploi[\s_]du[\s_]temps|agenda", filename, re.IGNORECASE))
        has_soutenance = bool(re.search(r"soutenances?|pfe", filename, re.IGNORECASE))
        return has_planning and has_soutenance

    @staticmethod
    def _is_student_list(columns: List[Column]) -> bool:
        """Liste d'étudiants reconnue sur les EN-TÊTES (Matricule + Nom), pas sur le nom du fichier."""
        leaves = [(c.leaf or "").lower() for c in columns]
        has_mat = any("matricule" in l or "inscription" in l for l in leaves)
        has_nom = any(re.search(r"\bnom(s)?\b", l) for l in leaves)
        return has_mat and has_nom

    # ------------------------------------------------------------------
    # Métadonnées académiques : bloc > (année : bandeau > feuille > fichier ; niveau/formation : fichier > feuille > bandeau)
    # ------------------------------------------------------------------
    @staticmethod
    def _extract(text: str) -> Dict[str, Any]:
        return AcademicMetadataExtractor.extract_from_text(text or "") or {}

    def _meta(self, ctx: SheetContext, local_text: str, cache: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
        if local_text not in cache:
            cache[local_text] = self._extract(local_text) if local_text else {}
        local = cache[local_text]

        def pick(key: str, order: List[Dict[str, Any]]):
            for src in order:
                if src.get(key):
                    return src[key]
            return None

        annee = pick("annee_universitaire", [local, ctx.pre_meta, ctx.sheet_meta, ctx.file_meta])
        if annee and "/" in annee:
            annee = annee.replace("/", "-")

        return {
            "annee_universitaire": annee,
            "niveau": pick("niveau", [local, ctx.file_meta, ctx.sheet_meta, ctx.pre_meta]),
            "formation": pick("formation", [local, ctx.file_meta, ctx.sheet_meta, ctx.pre_meta]),
        }

    @staticmethod
    def _header(ctx: SheetContext, *extra: Optional[str]) -> str:
        parts = [ctx.doc_title, ctx.sheet_display, *ctx.preamble, *extra]
        return "[" + " — ".join(dedupe_parts([p for p in parts if p])) + "]"

    @staticmethod
    def _chunk(text: str, doc: str, family: str, sheet: str, block_range: str,
               block_title: Optional[str], meta: Dict[str, Any]) -> Chunk:
        return Chunk(text=text, metadata=ChunkMetadata(
            document_name=doc, family_type=family, sheet_name=sheet,
            block_range=block_range, block_title=block_title, **meta))

    @staticmethod
    def _notes_line(ctx: SheetContext) -> List[str]:
        return [f"Remarque: {n}" for n in ctx.notes]

    # ------------------------------------------------------------------
    # Point d'entrée
    # ------------------------------------------------------------------
    def parse(self, file_path: str, document_title: Optional[str] = None,
              effective_start_line: int = 0) -> List[Chunk]:
        filename = document_title if document_title else os.path.basename(file_path)
        doc_title = re.sub(r"\s+", " ", re.sub(r"\.(xlsx|xlsm|xls)$", "", filename, flags=re.I).replace("_", " ")).strip()
        file_meta = self._extract(filename)
        planning_hint = self._is_planning_file(filename)
        wb = openpyxl.load_workbook(file_path, data_only=True)

        sheets = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                logger.info("Feuille masquée ignorée : %s", ws.title)

        chunks: List[Chunk] = []
        for idx, ws in enumerate(sheets):
            view = SheetView(ws)
            sheet_year = year_pair(ws.title)
            is_contact_sheet = bool(re.search(r"emails?|annuaire|contact", ws.title, re.I))
            ctx = SheetContext(
                doc_name=filename,
                doc_title=doc_title,
                sheet_name=ws.title,
                sheet_display=(sheet_year or ws.title) if len(sheets) > 1 else None,
                file_meta={} if is_contact_sheet else file_meta,
                sheet_meta=self._extract(sheet_year or ws.title),
                planning_hint=planning_hint,
            )
            start_row = effective_start_line if idx == 0 else 0
            layouts = [analyze_block(view, a, b) for a, b in view.blocks(start_row)]
            self._collect_context(view, layouts, ctx)

            cache: Dict[str, Dict[str, Any]] = {}
            state = TextState()
            before = len(chunks)
            for layout in layouts:
                if layout.kind == "text":
                    chunks += self._text_chunks(view, layout, ctx, state, cache)
                elif layout.kind == "table":
                    chunks += self._table_chunks(view, layout, ctx, cache)
            logger.info("%s / %s : %d chunks", filename, ws.title, len(chunks) - before)

        if not chunks:
            logger.warning("Aucun chunk produit pour %s", filename)
        return chunks

    # ------------------------------------------------------------------
    # Contexte de feuille : bandeaux, version, remarques (NB…)
    # ------------------------------------------------------------------
    def _collect_context(self, view: SheetView, layouts: List[BlockLayout], ctx: SheetContext) -> None:
        # Un petit bloc de texte AVANT le premier tableau est un en-tête (non fusionné), pas une liste de sujets
        first_table = next((i for i, l in enumerate(layouts) if l.kind == "table"), None)
        if first_table is not None:
            for l in layouts[:first_table]:
                if l.kind != "text":
                    continue
                lines = [view.raw(r, l.text_col) for r in range(l.r0, l.r1 + 1) if view.raw(r, l.text_col)]
                if len(lines) <= 6 and all(len(t) <= 120 and not (SECTION_RE.match(t) or SOURCE_RE.match(t)
                                                                 or NOTE_RE.match(t)) for t in lines):
                    l.kind, l.titles = "banner", lines

        def absorb(line: str) -> None:
            if NOTE_RE.match(line):
                note = NOTE_RE.sub("", line, count=1).strip()
                if note and note not in ctx.notes:
                    ctx.notes.append(note)
            elif VERSION_RE.match(line):
                ctx.version = ctx.version or line
            elif not INSTITUTION_RE.match(line) and line not in ctx.preamble:
                ctx.preamble.append(line)

        for l in layouts:
            if l.kind == "banner":
                for t in l.titles:
                    absorb(t)
            elif l.kind == "table":
                # titres propres au tableau : on n'en extrait que version et remarques ici
                for t in l.titles:
                    if NOTE_RE.match(t) or VERSION_RE.match(t):
                        absorb(t)
            elif l.kind == "text":
                for r in range(l.r0, l.r1 + 1):
                    t = view.raw(r, l.text_col)
                    if t and NOTE_RE.match(t):
                        absorb(t)
        ctx.pre_meta = self._extract(" ".join(ctx.preamble))

    # ------------------------------------------------------------------
    # Texte libre (C-texte-libre) : l'état (section, source) traverse les blocs de la feuille
    # ------------------------------------------------------------------
    def _text_chunks(self, view: SheetView, layout: BlockLayout, ctx: SheetContext,
                     state: TextState, cache: Dict[str, Dict[str, Any]]) -> List[Chunk]:
        out: List[Chunk] = []
        c = layout.text_col
        for r in range(layout.r0, layout.r1 + 1):
            text = view.raw(r, c)
            if not text or NOTE_RE.match(text):
                continue
            if SECTION_RE.match(text):
                state.section, state.source = text.rstrip(" :"), None
                continue
            m = SOURCE_RE.match(text)
            if m:
                state.source = m.group(1).strip() or None
                continue
            for item in split_items(view, r, c):
                label = "Sujet" if re.search(r"sujet", " ".join(filter(None, [state.section, ctx.doc_title])), re.I) \
                    else "Élément"
                lines = [self._header(ctx, state.section)]
                if state.source:
                    lines.append(f"Source: {state.source}")
                lines.append(f"{label}: {item}")
                lines += self._notes_line(ctx)
                out.append(self._chunk("\n".join(lines), ctx.doc_name,
                                       "C-texte-libre", ctx.sheet_name, f"{layout.r0}-{layout.r1}",
                                       state.section, self._meta(ctx, state.section or "", cache)))
        return out

    # ------------------------------------------------------------------
    # Tableaux
    # ------------------------------------------------------------------
    def _table_chunks(self, view: SheetView, layout: BlockLayout, ctx: SheetContext,
                      cache: Dict[str, Dict[str, Any]]) -> List[Chunk]:
        body = layout.body_rows
        is_student = self._is_student_list(layout.columns)
        fill = {c.idx: sum(1 for r in body if view.body_value(r, c.idx) is not None) for c in layout.columns}
        kept = [c for c in layout.columns
                if fill[c.idx] > 0 and not (is_student and any(ATTENDANCE_RE.match(p) for p in c.path))]
        if not kept or not body:
            return []

        groups: "OrderedDict[Tuple[str, ...], List[Column]]" = OrderedDict()
        for c in kept:
            if len(c.path) >= 2:
                groups.setdefault(tuple(c.path[:-1]), []).append(c)
        signatures = Counter(tuple(col.leaf or "" for col in cols) for cols in groups.values())
        if len(groups) >= 2 and signatures and signatures.most_common(1)[0][1] >= 2 \
                and len(next(iter(groups.values()))) >= 2:
            return self._grid_chunks(view, layout, ctx, kept, groups, cache)
        return self._flat_chunks(view, layout, ctx, kept, fill, is_student, cache)

    # --- 4a. Grille salles x créneaux (famille D) -----------------------------------------------
    def _grid_chunks(self, view: SheetView, layout: BlockLayout, ctx: SheetContext, kept: List[Column],
                     groups: "OrderedDict[Tuple[str, ...], List[Column]]",
                     cache: Dict[str, Dict[str, Any]]) -> List[Chunk]:
        in_group = {c.idx for cols in groups.values() for c in cols}
        first_g = min(in_group)
        key_cols = [c for c in kept if c.idx not in in_group and c.idx < first_g]
        extra_cols = [c for c in kept if c.idx not in in_group and c.idx > first_g]

        def role(label: str) -> Optional[str]:
            if re.search(r"pr[ée]sident", label, re.I):
                return "president"
            if re.search(r"membre|examinat|jury", label, re.I):
                return "membres"
            if re.search(r"projet|pfe|sujet|n°", label, re.I):
                return "projet"
            return None

        leaf_roles = {role(c.leaf or "") for cols in groups.values() for c in cols}
        soutenance = "projet" in leaf_roles and bool(leaf_roles & {"president", "membres"})

        def key_label(label: str) -> str:
            if not soutenance:
                return label
            if re.search(r"journ|date|jour", label, re.I):
                return "Date"
            if re.search(r"cr[ée]neau|horaire|heure", label, re.I):
                return "Créneau"
            return label

        # colonne sans titre -> rattachée au bandeau de tête situé immédiatement à sa gauche
        extra_top: Dict[int, Optional[str]] = {}
        for e in extra_cols:
            left = [(max(c.idx for c in cols), key) for key, cols in groups.items() if max(c.idx for c in cols) < e.idx]
            nearest = max(left)[1] if left else ()
            extra_top[e.idx] = nearest[0] if len(nearest) >= 2 else (nearest[-1] if nearest else None)

        # Famille D = grille de planning (soutenances, ou lignes Date/Jour/Créneau/Heure) ; sinon simple tableau C
        schedule = soutenance or ctx.planning_hint or any(
            re.search(r"journ|date|jour|cr[ée]neau|heure|horaire", c.path[-1] if c.path else "", re.I) for c in key_cols)
        family = "D" if schedule else "C"

        out: List[Chunk] = []
        band_title = " — ".join(layout.titles) if layout.titles else None
        for r in layout.body_rows:
            keys = []
            for c in key_cols:
                v = view.body_value(r, c.idx)
                if v:
                    keys.append((key_label(c.path[-1] if c.path else f"Colonne {get_column_letter(c.idx)}"), v))
            for key, cols in groups.items():
                vals = [(c.leaf or "", view.body_value(r, c.idx)) for c in cols]
                if not any(v for _, v in vals):
                    continue
                salle, top = key[-1], (key[0] if len(key) >= 2 else None)
                notes: List[str] = []
                slot_override = None
                for e in extra_cols:
                    ev = view.body_value(r, e.idx)
                    match_top = (extra_top.get(e.idx) == top) if top is not None else (extra_top.get(e.idx) == salle)
                    if ev and match_top:
                        if TIME_RANGE_RE.match(ev):
                            slot_override = ev
                        else:
                            notes.append(ev)

                lines = []
                if soutenance:
                    k_line = [("Salle", salle)] + [(k, v) for k, v in keys]
                    if slot_override:
                        k_line = [(k, slot_override if k == "Créneau" else v) for k, v in k_line]
                        if not any(k == "Créneau" for k, _ in k_line):
                            k_line.append(("Créneau", slot_override))
                    d_line = []
                    for label, v in vals:
                        if not v:
                            continue
                        rl = role(label)
                        if rl == "president":
                            d_line.append(("Président", "; ".join(split_people(v)) or v))
                        elif rl == "membres":
                            people = "; ".join(split_people(v))
                            if people:
                                d_line.append(("Membres", people))
                        elif rl == "projet":
                            d_line.append(("Projet", v))
                        else:
                            d_line.append((label, v))
                    d_line.sort(key=lambda kv: {"Projet": 0, "Président": 1, "Membres": 2}.get(kv[0], 3))
                    lines.append("Soutenance PFE " + self._header(ctx, band_title, top))
                    lines.append(" | ".join(f"{k}: {v}" for k, v in k_line))
                    lines.append(" | ".join(f"{k}: {v}" for k, v in d_line + [("Remarque", n) for n in notes]))
                else:
                    lines.append(self._header(ctx, band_title, top))
                    lines.append(" | ".join([f"Groupe: {salle}"] + [f"{k}: {v}" for k, v in keys]))
                    extra_items = []
                    for e in extra_cols:
                        ev = view.body_value(r, e.idx)
                        match_top = (extra_top.get(e.idx) == top) if top is not None else (extra_top.get(e.idx) == salle)
                        if ev and match_top:
                            col_name = e.leaf or (e.path[-1] if e.path else "Remarque")
                            extra_items.append(f"{col_name}: {ev}")
                    lines.append(" | ".join([f"{l}: {v}" for l, v in vals if v] + extra_items
                                            + [f"Remarque: {n}" for n in notes]))
                lines += self._notes_line(ctx)
                title_for_meta = " ".join(filter(None, [band_title, top]))
                out.append(self._chunk("\n".join(lines), ctx.doc_name, family, ctx.sheet_name,
                                       f"{layout.table_r0}-{layout.r1}", top or band_title,
                                       self._meta(ctx, title_for_meta, cache)))
        return out

    # --- 4b. Tableau plat / liste d'étudiants (famille C ; D si indice 'planning' dans le nom) ---
    def _flat_chunks(self, view: SheetView, layout: BlockLayout, ctx: SheetContext, kept: List[Column],
                     fill: Dict[int, int], is_student: bool, cache: Dict[str, Dict[str, Any]]) -> List[Chunk]:
        body = layout.body_rows
        section: Optional[str] = None
        labels: List[str] = []
        for c in kept:
            label = " / ".join(c.path)
            if not label:   # colonne sans titre : remarque ponctuelle ou colonne de données
                label = "Remarque" if fill[c.idx] <= max(1, 0.2 * len(body)) else f"Colonne {get_column_letter(c.idx)}"
            m = HEADER_SECTION_RE.match(label)
            if m:           # en-tête du type 'Sujets Licence' : colonne = Sujet, section = niveau
                section, label = section or label, "Sujet"
            base, n = label, 2
            while label in labels:
                label, n = f"{base} ({n})", n + 1
            labels.append(label)

        rows: List[Tuple[int, "OrderedDict[str, str]"]] = []
        for r in body:
            # Sauter les lignes d'en-tête répétées dans le corps (ex. en-têtes de saut de page)
            if all(view.raw(r, c.idx) == c.leaf for c in kept if c.leaf and view.raw(r, c.idx)):
                continue

            # Absorber les lignes de notes en bas de tableau plutôt que de les indexer comme un enregistrement
            row_vals = [view.body_value(r, c.idx) for c in kept if view.body_value(r, c.idx)]
            if len(row_vals) == 1 and NOTE_RE.match(str(row_vals[0])):
                note = NOTE_RE.sub("", str(row_vals[0]), count=1).strip()
                if note and note not in ctx.notes:
                    ctx.notes.append(note)
                continue

            rec: "OrderedDict[str, str]" = OrderedDict()
            for c, label in zip(kept, labels):
                v = view.body_value(r, c.idx)
                if v:
                    rec[label] = v
            if rec:
                rows.append((r, rec))
        if not rows:
            return []

        # Colonnes constantes (ex. Palier/Spécialité/Section d'une liste) : remontées dans l'en-tête
        constants: "OrderedDict[str, str]" = OrderedDict()
        if len(rows) >= 5:
            for label in labels:
                vals = {rec.get(label) for _, rec in rows}
                if len(vals) == 1 and None not in vals:
                    constants[label] = vals.pop()
        const_line = " | ".join(f"{k}: {v}" for k, v in constants.items())

        titles = [t for t in layout.titles if not (VERSION_RE.match(t) or NOTE_RE.match(t) or INSTITUTION_RE.match(t))]
        extra_title = " — ".join(titles) if titles else None
        header = self._header(ctx, extra_title, section)
        family = "D" if ctx.planning_hint else "C"
        meta = self._meta(ctx, " ".join(filter(None, [extra_title, section])), cache)
        block_range = f"{layout.table_r0}-{layout.r1}"
        block_title = section or extra_title

        out: List[Chunk] = []
        for _, rec in rows:
            fields = [(k, v) for k, v in rec.items() if k not in constants]
            if not fields:
                continue
            lines = [header] + ([const_line] if const_line else [])
            lines.append(" | ".join(f"{k}: {v}" for k, v in fields))
            lines += self._notes_line(ctx)
            out.append(self._chunk("\n".join(lines), ctx.doc_name, family, ctx.sheet_name,
                                   block_range, block_title, meta))

        if is_student:
            out += self._student_summary(ctx, header, const_line, labels, constants, rows, family,
                                         block_range, block_title, meta)
        return out

    # --- 4c. Liste d'étudiants : effectif + listes par groupe (unité naturelle d'un affichage) ---
    def _student_summary(self, ctx: SheetContext, header: str, const_line: str, labels: List[str],
                         constants: Dict[str, str], rows: List[Tuple[int, Dict[str, str]]], family: str,
                         block_range: str, block_title: Optional[str], meta: Dict[str, Any]) -> List[Chunk]:
        recs = [rec for _, rec in rows]
        n = len(recs)
        free = [l for l in labels if l not in constants]
        categorical = []
        for l in free:
            counts = Counter(rec.get(l) for rec in recs if rec.get(l))
            if 2 <= len(counts) <= 12 and len(counts) < n:
                categorical.append((l, counts))

        base = [header] + ([const_line] if const_line else [])
        over = [f"Effectif: {n} étudiants" + (f" ({ctx.version})" if ctx.version else "")]
        for l, counts in categorical:
            over.append(f"{l}: " + ", ".join(f"{v} ({counts[v]})" for v in sorted(counts, key=_natural_key)))
        out = [self._chunk("\n".join(base + over + self._notes_line(ctx)),
                           ctx.doc_name, family, ctx.sheet_name, block_range, "Effectif", meta)]

        group_label = next((l for l, _ in categorical if re.search(r"group", l, re.I)), None)
        roster_cols = [l for l in free if l != group_label]
        if group_label:
            parts = OrderedDict()
            for v in sorted({rec[group_label] for rec in recs if rec.get(group_label)}, key=_natural_key):
                parts[f"{group_label}: {v}"] = [rec for rec in recs if rec.get(group_label) == v]
        else:
            parts = OrderedDict((f"Liste (partie {i // 40 + 1})", recs[i:i + 40]) for i in range(0, n, 40))
        for title, members in parts.items():
            for i in range(0, len(members), 40):
                chunk_members = members[i:i + 40]
                suffix = f" (partie {i // 40 + 1})" if len(members) > 40 else ""
                lines = base + [f"{title}{suffix} — {len(members)} étudiants", " | ".join(roster_cols)]
                lines += [" | ".join(rec.get(l, "") for l in roster_cols) for rec in chunk_members]
                out.append(self._chunk("\n".join(lines + self._notes_line(ctx)), ctx.doc_name, family,
                                       ctx.sheet_name, block_range, title, meta))
        return out